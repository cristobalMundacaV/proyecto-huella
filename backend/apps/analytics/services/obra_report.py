"""CARBONO-ZERO-V1 — Informe Ambiental de Obra (PDF + Excel).

Builds directly on `obra_environmental_dashboard.build_obra_dashboard` —
the single motor — so every figure in the report is exactly the figure
already shown on the dashboard and used by the AI copiloto. This module
only renders that same payload into two downloadable formats; it computes
nothing new.

Deliberately NOT wired into `InformeAmbiental`/`generate_report`
(the professional-review governance report pipeline in
`services/professional_v2.py`): that pipeline's immutability/validation
workflow is designed around dossier-type governance artifacts
(evidencia/cálculo/problemática/expediente), not a routine period report a
professional generates and re-generates freely. This is a separate,
lighter "generate & download" capability — a deliberate scope choice, not
an oversight.
"""
import io

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .obra_environmental_dashboard import build_obra_dashboard

PAGE_MARGIN = 45
LINE_HEIGHT = 13
WRAP_WIDTH = 105


def _fmt(value, unit=""):
    if value is None:
        return "sin dato"
    if unit:
        return f"{value} {unit}"
    return str(value)


def _flow_line(label, unit_stats):
    if not unit_stats:
        return f"{label}: sin datos en el período."
    parts = []
    for unit, stats in unit_stats.items():
        parts.append(f"{stats['total']} {unit} (promedio {stats['promedio']} {unit}/mes)")
    return f"{label}: " + "; ".join(parts)


def build_report_sections(dashboard):
    """Plain-text section list — shared by the PDF renderer and, if a
    future format needs it, anything else. Never invents a figure not
    already present in `dashboard`."""
    kpis = dashboard["kpis"]
    readiness = dashboard["readiness"]
    risk = dashboard["risk"]
    resumen = dashboard["resumen_ejecutivo"]

    sections = []

    sections.append(("1. Identificación de la obra", [
        f"Organización: {dashboard['organizacion_nombre']} ({dashboard['organizacion_id']})",
        f"Obra: {dashboard['obra_nombre']} (id {dashboard['obra_id']})",
    ]))
    sections.append(("2. Período y alcance", [
        f"Período actual: {dashboard['period']['start']} a {dashboard['period']['end']}",
        f"Período de comparación: {dashboard['previous_period']['start']} a {dashboard['previous_period']['end']}",
        f"Versión de reglas de diagnóstico: {dashboard['ruleset_version']}",
    ]))
    sections.append(("3. Resumen ejecutivo", [
        f"Estado ambiental: {dashboard['estado_ejecutivo']['label']}",
        f"Nivel de riesgo ambiental-operacional: {risk['nivel']} ({risk['risk_score']}/100)",
        f"Hallazgos de alta severidad: {resumen['hallazgos_altos']}",
        f"Evidencias faltantes detectadas: {resumen['evidencias_faltantes']}",
        f"Factores ambientales pendientes de asignar: {resumen['factores_pendientes']}",
        f"Cobertura de registros del período: {_fmt(resumen['cobertura_periodo_pct'], '%')}",
    ]))
    sections.append(("4. Estado ambiental general", [
        f"Clasificación determinista: {dashboard['estado_ejecutivo']['label']} "
        f"(calculada por reglas de riesgo/cobertura, nunca por criterio del asistente de IA).",
    ]))
    sections.append(("5. Inventario GEI (gases de efecto invernadero)", [
        f"Huella total: {_fmt(kpis['huella_total_tco2e'], 'tCO2e')}",
        f"Alcance 1 (combustión directa): {_fmt(kpis['alcance_1_tco2e'], 'tCO2e')}",
        f"Alcance 2 (energía comprada): {_fmt(kpis['alcance_2_tco2e'], 'tCO2e')}",
        f"Alcance 3 (materiales, residuos, agua y otras fuentes indirectas): {_fmt(kpis['alcance_3_tco2e'], 'tCO2e')}",
    ]))
    sections.append(("6. Combustibles", [_flow_line("Consumo de combustible", kpis["combustible"])]))
    sections.append(("7. Energía", [_flow_line("Consumo de energía", kpis["energia"])]))
    sections.append(("8. Agua", [_flow_line("Consumo de agua", kpis["agua"])]))
    sections.append(("9. Residuos", [
        _flow_line("Generación de residuos", kpis["residuos"]),
        f"Tasa de valorización: {_fmt(kpis['tasa_valorizacion_pct'], '%')}",
    ]))
    sections.append(("10. Materiales", [_flow_line("Consumo/impacto de materiales", kpis["materiales"])]))
    sections.append(("11. Transporte", ["Sin datos de transporte modelados para esta obra en este período."]))
    sections.append(("12. Maquinaria", [
        finding["title"] for finding in dashboard["top_findings"] if finding["entity_type"] == "activo"
    ] or ["Sin hallazgos de maquinaria en este período."]))
    sections.append(("13. Intensidades", [
        "Las intensidades por unidad de actividad (m2, m3, unidad de producción) requieren una línea base "
        "de actividad configurada; consulte los indicadores gobernados de la obra para el detalle vigente.",
    ]))
    sections.append(("14. Calidad de datos", [
        f"Cobertura de evidencia: {_fmt(readiness['evidencia_pct'], '%')}",
        f"Cobertura de factores ambientales asignados: {_fmt(readiness['factores_pct'], '%')}",
    ]))
    sections.append(("15. Evidencias", [
        f"Pendientes de evidencia: {resumen['evidencias_faltantes']} hallazgo(s) de evidencia incompleta.",
    ]))
    sections.append(("16. Hallazgos", [
        f"[{finding['severity'].upper()}] {finding['title']}" for finding in dashboard["top_findings"]
    ] or ["Sin hallazgos en este período."]))
    sections.append(("17. Riesgos", [
        f"Riesgo global: {risk['nivel']} ({risk['risk_score']}/100)",
        *[f"  - {categoria}: {detalle['findings']} hallazgo(s), score {detalle['score']}" for categoria, detalle in risk["por_categoria"].items()],
    ]))
    sections.append(("18. Acciones", [
        f"  - {recomendacion}"
        for finding in dashboard["top_findings"] for recomendacion in finding.get("recommendations", [])
    ][:15] or ["Sin acciones pendientes registradas."]))
    sections.append(("19. Factores y metodología", [
        "Los factores ambientales aplicados provienen del catálogo gobernado de Carbono Zero "
        "(EC3/ÖKOBAUDAT y factores propios aprobados) — ver trazabilidad para el detalle por cálculo.",
    ]))
    sections.append(("20. Trazabilidad", [
        "Cada cifra de este informe puede reconstruirse mediante trace_metric_provenance / "
        "get_factor_provenance hasta el registro, evidencia y factor de origen.",
    ]))
    sections.append(("21. Anexos", [
        f"Readiness del período: {'LISTO PARA REPORTE' if readiness['listo_para_reporte'] else 'NO LISTO PARA CIERRE'}",
        *[f"  - Pendiente: {item}" for item in readiness["pendientes"]],
    ]))
    return sections


def render_report_pdf(dashboard):
    sections = build_report_sections(dashboard)
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4, pageCompression=0)
    width, height = A4
    y = height - 55

    def write(text, *, bold=False):
        nonlocal y
        chunks = [text[index:index + WRAP_WIDTH] for index in range(0, len(text), WRAP_WIDTH)] or [""]
        for chunk in chunks:
            if y < 50:
                pdf.showPage()
                y = height - 55
            pdf.setFont("Helvetica-Bold" if bold else "Helvetica", 10)
            pdf.drawString(PAGE_MARGIN, y, chunk)
            y -= LINE_HEIGHT

    write("CARBONO ZERO — INFORME AMBIENTAL DE OBRA", bold=True)
    write(f"Obra: {dashboard['obra_nombre']} | Organización: {dashboard['organizacion_nombre']}")
    write(f"Período: {dashboard['period']['start']} a {dashboard['period']['end']}")
    write("")
    for title, lines in sections:
        write(title, bold=True)
        for line in lines:
            write(str(line))
        write("")
    pdf.save()
    return output.getvalue()


def render_report_excel(dashboard):
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Informe Ambiental de Obra"])
    summary.append(["Obra", dashboard["obra_nombre"]])
    summary.append(["Organización", dashboard["organizacion_nombre"]])
    summary.append(["Período", f"{dashboard['period']['start']} a {dashboard['period']['end']}"])
    summary.append(["Estado ambiental", dashboard["estado_ejecutivo"]["label"]])
    summary.append(["Riesgo", f"{dashboard['risk']['nivel']} ({dashboard['risk']['risk_score']}/100)"])
    summary.append([])
    summary.append(["KPI", "Valor", "Unidad"])
    kpis = dashboard["kpis"]
    summary.append(["Huella total", kpis["huella_total_tco2e"], "tCO2e"])
    summary.append(["Alcance 1", kpis["alcance_1_tco2e"], "tCO2e"])
    summary.append(["Alcance 2", kpis["alcance_2_tco2e"], "tCO2e"])
    summary.append(["Alcance 3", kpis["alcance_3_tco2e"], "tCO2e"])
    summary.append(["Tasa de valorización", kpis["tasa_valorizacion_pct"], "%"])
    summary.append(["Cobertura de evidencia", kpis["cobertura_evidencia_pct"], "%"])

    flows = workbook.create_sheet("Flujos")
    flows.append(["Métrica", "Unidad", "Total", "Promedio", "Máximo", "Mínimo", "Períodos con datos"])
    for metric in ("agua", "combustible", "energia", "residuos", "materiales"):
        for unit, stats in (kpis.get(metric) or {}).items():
            flows.append([metric, unit, stats["total"], stats["promedio"], stats["maximo"], stats["minimo"], stats["periodos_con_datos"]])

    findings = workbook.create_sheet("Hallazgos")
    findings.append(["Código", "Severidad", "Prioridad", "Título", "Métrica", "Entidad"])
    for finding in dashboard["top_findings"]:
        findings.append([finding["code"], finding["severity"], finding["priority_score"], finding["title"], finding.get("metric"), finding.get("entity_name")])

    readiness_sheet = workbook.create_sheet("Readiness")
    readiness = dashboard["readiness"]
    readiness_sheet.append(["Indicador", "Valor (%)"])
    readiness_sheet.append(["Cobertura de registros", readiness["cobertura_registros_pct"]])
    readiness_sheet.append(["Evidencia", readiness["evidencia_pct"]])
    readiness_sheet.append(["Factores", readiness["factores_pct"]])
    readiness_sheet.append(["Validación profesional", readiness["validacion_profesional_pct"]])
    readiness_sheet.append([])
    readiness_sheet.append(["Listo para reporte", "SÍ" if readiness["listo_para_reporte"] else "NO"])
    readiness_sheet.append(["Pendientes"])
    for item in readiness["pendientes"]:
        readiness_sheet.append([item])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_obra_report(organizacion, user, obra, *, date_from=None, date_to=None, relative_months=None):
    """Convenience: dashboard + both renders in one call, for the view."""
    dashboard = build_obra_dashboard(organizacion, user, obra, date_from=date_from, date_to=date_to, relative_months=relative_months)
    return {
        "dashboard": dashboard,
        "pdf": render_report_pdf(dashboard),
        "excel": render_report_excel(dashboard),
    }
