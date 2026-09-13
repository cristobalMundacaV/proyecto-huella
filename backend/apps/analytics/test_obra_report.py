"""CARBONO-ZERO-V1 — Informe Ambiental de Obra (PDF + Excel).

Verifies the report renderers never recompute a number themselves (every
figure must trace back to `build_obra_dashboard`'s own output) and that
both file formats are real, well-formed, downloadable artifacts.
"""
from datetime import date

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from apps.analytics.models import Obra, Organizacion, UsuarioOrganizacion
from apps.analytics.services.obra_environmental_dashboard import build_obra_dashboard
from apps.analytics.services.obra_report import build_report_sections, render_report_excel, render_report_pdf

User = get_user_model()


class ObraReportDemoTenantTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_ai_demo_tenant")
        cls.org = Organizacion.objects.get(organizacion_id="DEMO_HORIZONTE")
        membership = UsuarioOrganizacion.objects.filter(organizacion=cls.org).select_related("user").first()
        cls.admin = membership.user
        cls.obra_norte = Obra.objects.get(organizacion=cls.org, nombre="Edificio Horizonte Norte")
        cls.dashboard = build_obra_dashboard(cls.org, cls.admin, cls.obra_norte, date_from="2020-01-01", date_to="2035-12-31")

    def test_sections_cite_the_same_huella_total_as_the_dashboard(self):
        sections = build_report_sections(self.dashboard)
        gei_section = dict(sections)["5. Inventario GEI (gases de efecto invernadero)"]
        expected = f"{self.dashboard['kpis']['huella_total_tco2e']} tCO2e"
        self.assertTrue(any(expected in line for line in gei_section), gei_section)

    def test_sections_cite_the_same_risk_level_as_the_dashboard(self):
        sections = build_report_sections(self.dashboard)
        resumen = dict(sections)["3. Resumen ejecutivo"]
        expected = self.dashboard["risk"]["nivel"]
        self.assertTrue(any(expected in line for line in resumen))

    def test_report_has_all_21_sections(self):
        sections = build_report_sections(self.dashboard)
        self.assertEqual(len(sections), 21)
        self.assertTrue(sections[0][0].startswith("1."))
        self.assertTrue(sections[-1][0].startswith("21."))

    def test_pdf_is_a_real_non_trivial_pdf(self):
        content = render_report_pdf(self.dashboard)
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 1000)

    def test_excel_is_a_real_loadable_workbook_with_matching_totals(self):
        content = render_report_excel(self.dashboard)
        workbook = load_workbook(filename=__import__("io").BytesIO(content))
        self.assertIn("Resumen", workbook.sheetnames)
        self.assertIn("Flujos", workbook.sheetnames)
        self.assertIn("Hallazgos", workbook.sheetnames)
        self.assertIn("Readiness", workbook.sheetnames)
        summary_rows = list(workbook["Resumen"].iter_rows(values_only=True))
        huella_row = next(row for row in summary_rows if row[0] == "Huella total")
        self.assertEqual(huella_row[1], self.dashboard["kpis"]["huella_total_tco2e"])

    def test_never_invents_a_number_outside_the_dashboard_payload(self):
        # Every numeric KPI referenced in the PDF sections must originate
        # from the dashboard dict — cross-check a sample of them.
        sections = dict(build_report_sections(self.dashboard))
        kpis = self.dashboard["kpis"]
        self.assertIn(str(kpis["alcance_1_tco2e"]), "\n".join(sections["5. Inventario GEI (gases de efecto invernadero)"]))
        self.assertIn(str(kpis["alcance_2_tco2e"]), "\n".join(sections["5. Inventario GEI (gases de efecto invernadero)"]))
        self.assertIn(str(kpis["alcance_3_tco2e"]), "\n".join(sections["5. Inventario GEI (gases de efecto invernadero)"]))


class ObraReportApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_ai_demo_tenant")
        cls.org = Organizacion.objects.get(organizacion_id="DEMO_HORIZONTE")
        membership = UsuarioOrganizacion.objects.filter(organizacion=cls.org).select_related("user").first()
        cls.admin = membership.user
        cls.obra_norte = Obra.objects.get(organizacion=cls.org, nombre="Edificio Horizonte Norte")

    def test_pdf_endpoint_downloads_a_real_pdf(self):
        self.client.force_login(self.admin)
        url = f"/api/organizaciones/{self.org.organizacion_id}/obras/{self.obra_norte.id}/informe-ambiental.pdf"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_excel_endpoint_downloads_a_real_workbook(self):
        self.client.force_login(self.admin)
        url = f"/api/organizaciones/{self.org.organizacion_id}/obras/{self.obra_norte.id}/informe-ambiental.xlsx"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])

    def test_report_endpoints_require_authentication(self):
        url = f"/api/organizaciones/{self.org.organizacion_id}/obras/{self.obra_norte.id}/informe-ambiental.pdf"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 401)

    def test_report_endpoints_never_leak_a_foreign_tenants_obra(self):
        other_org = Organizacion.objects.create(nombre="Report org B")
        foreign_obra = Obra.objects.create(organizacion=other_org, nombre="Obra ajena informe", fecha_inicio=date(2026, 1, 1))
        self.client.force_login(self.admin)
        url = f"/api/organizaciones/{self.org.organizacion_id}/obras/{foreign_obra.id}/informe-ambiental.pdf"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 404)
